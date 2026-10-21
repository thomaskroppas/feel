"""Выгрузка xlsx: три листа, кириллица, даты ячейками Excel."""
import io
from datetime import date, time

from openpyxl import load_workbook


def test_export_xlsx(client_for):
    a = client_for("a@example.com")
    b = client_for("b@example.com")
    tags = a.get("/api/tags").json()
    friends = next(t for t in tags if t["name"] == "Друзья")
    home = next(t for t in tags if t["name"] == "Дома")
    a.post("/api/checkins", {"emotion_id": "radost", "occurred_at": "2026-10-19T23:30",
                             "note": "Вечер с друзьями", "tag_ids": [friends["id"], home["id"]]})
    a.put("/api/days/2026-10-19", {"plan_status": "yes", "plan_text": "Отдых", "sleep_hours": 7.5,
                                   "sleep_quality": 4, "sport_type": "", "day_summary": "Хорошо"})
    b.post("/api/checkins", {"emotion_id": "trevoga", "note": "чужое"})

    r = a.get("/api/export.xlsx?period=all")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Записи", "Дни", "Метки"]

    ws = wb["Записи"]
    assert [c.value for c in ws[1]][:6] == ["Дата", "Время", "Слово", "Квадрант", "Энергия", "Приятность"]
    row = [c for c in ws[2]]
    assert ws.max_row == 2  # запись пользователя Б не попала
    assert row[0].is_date and row[0].value.date() == date(2026, 10, 19)  # местная дата, не UTC
    assert row[1].value == time(23, 30)
    assert row[2].value == "Радость" and row[3].value == "Высокая энергия · приятно"
    assert row[4].value == 9 and row[5].value == 11
    assert row[6].value == "Друзья" and row[7].value == "Дома" and row[9].value == "Вечер с друзьями"

    days = wb["Дни"]
    d = [c.value for c in days[2]]
    assert d[1:] == ["есть", "Отдых", 7.5, 4, "не было", None, "Хорошо"]
    assert days["A2"].is_date

    tags_ws = {(r[0].value, r[1].value): r[2].value for r in wb["Метки"].iter_rows(min_row=2)}
    assert tags_ws[("С кем", "Друзья")] == 1 and tags_ws[("Где", "Дома")] == 1
    assert tags_ws[("Чем занят", "Работа")] == 0


def test_export_requires_login(api):
    assert api.get("/api/export.xlsx").status_code == 401
