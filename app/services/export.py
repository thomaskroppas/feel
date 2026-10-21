"""Выгрузка в xlsx: записи, дни, метки. Даты и время — по часовому поясу пользователя,
ячейками формата даты Excel, а не текстом."""
import io
from collections import Counter
from datetime import date
from pathlib import PurePosixPath

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import models
from ..timeutil import to_local
from .days import checkins_between, day_logs_between

QUADRANT_TITLES = {"red": "Высокая энергия · неприятно", "yellow": "Высокая энергия · приятно",
                   "blue": "Низкая энергия · неприятно", "green": "Низкая энергия · приятно"}
GROUP_TITLES = {"with": "С кем", "where": "Где", "doing": "Чем занят"}
DATE_FMT = "DD.MM.YYYY"
TIME_FMT = "HH:MM"


def _sheet(wb: Workbook, title: str, headers: list, widths: list, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def build(db: Session, user: models.User, tz_name: str, start: date, end: date) -> bytes:
    checkins = checkins_between(db, user.id, start, end)
    logs = day_logs_between(db, user.id, start, end)
    wb = Workbook()
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = _sheet(wb, "Записи",
                ["Дата", "Время", "Слово", "Квадрант", "Энергия", "Приятность",
                 "С кем", "Где", "Чем занят", "Заметка", "Медиа"],
                [12, 8, 20, 28, 9, 11, 22, 22, 26, 50, 30], first=True)
    for c in checkins:
        local = to_local(c.occurred_at, tz_name)
        by_group = {g: ", ".join(t.name for t in c.tags if t.group == g) for g in GROUP_TITLES}
        ws.append([local.date(), local.time().replace(second=0, microsecond=0), c.emotion.name,
                   QUADRANT_TITLES[c.emotion.quadrant], c.emotion.energy, c.emotion.pleasantness,
                   by_group["with"], by_group["where"], by_group["doing"], c.note or "",
                   ", ".join(PurePosixPath(m.path).name for m in c.media)])
        row = ws.max_row
        ws.cell(row=row, column=1).number_format = DATE_FMT
        ws.cell(row=row, column=2).number_format = TIME_FMT
        ws.cell(row=row, column=10).alignment = wrap

    ws = _sheet(wb, "Дни",
                ["Дата", "План", "Текст плана", "Сон, ч", "Качество сна", "Спорт", "Минуты", "Итоги дня"],
                [12, 8, 40, 8, 13, 16, 9, 50])
    for d in sorted(logs):
        log = logs[d]
        plan = {"yes": "есть", "no": "нет"}.get(log.plan_status or "", "")
        sport = "" if log.sport_type is None else (log.sport_type or "не было")
        ws.append([d, plan, log.plan_text or "", log.sleep_hours, log.sleep_quality, sport,
                   log.sport_minutes, log.day_summary or ""])
        row = ws.max_row
        ws.cell(row=row, column=1).number_format = DATE_FMT
        ws.cell(row=row, column=3).alignment = wrap
        ws.cell(row=row, column=8).alignment = wrap

    ws = _sheet(wb, "Метки", ["Группа", "Метка", "Записей"], [14, 26, 10])
    used = Counter(t.id for c in checkins for t in c.tags)
    tags = db.scalars(select(models.Tag).where(models.Tag.user_id == user.id)
                      .order_by(models.Tag.group, models.Tag.sort, models.Tag.id))
    order = list(GROUP_TITLES)
    for t in sorted(tags, key=lambda t: (order.index(t.group), t.sort, t.id)):
        ws.append([GROUP_TITLES[t.group], t.name, used.get(t.id, 0)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
